# -*- coding: utf-8 -*-
"""D3-T2 整改通知单 + 复查意见测试。"""

import tempfile
from pathlib import Path

import pytest


@pytest.fixture(autouse=True)
def _reset():
    from src.baa_engine.report_generator.template import reset_template_registry_for_test
    reset_template_registry_for_test()
    yield
    reset_template_registry_for_test()


def test_build_correction_notice_docx(tmp_path):
    from src.baa_engine.word_generator import build_correction_notice_docx

    items = [
        {"id": "f1", "review_id": "r1", "function_id": "DIM-001", "entity_id": "e1", "status": "confirmed", "note": "净高不足", "reason": ""},
        {"id": "f2", "review_id": "r1", "function_id": "DIM-001", "entity_id": "e2", "status": "confirmed", "note": "净高不足", "reason": ""},
    ]
    meta = {"project_name": "测试项目", "drawing_name": "test.dxf", "reviewer": "张三"}
    out = build_correction_notice_docx(items, meta, str(tmp_path / "notice.docx"))
    assert Path(out).is_file()
    assert Path(out).suffix == ".docx"
    assert Path(out).stat().st_size > 0


def test_build_recheck_opinion_docx(tmp_path):
    from src.baa_engine.word_generator import build_recheck_opinion_docx

    rechecks = [
        {"function_id": "DIM-001", "entity_id": "e1", "status": "passed", "opinion": "已符合"},
        {"function_id": "DIM-001", "entity_id": "e2", "status": "failed", "opinion": "厚度仍不足"},
    ]
    meta = {"project_name": "测试项目", "drawing_name": "test.dxf", "reviewer": "李四", "corrector": "王五"}
    out = build_recheck_opinion_docx(rechecks, meta, str(tmp_path / "recheck.docx"))
    assert Path(out).is_file()
    assert Path(out).stat().st_size > 0


def test_docx_content_readable(tmp_path):
    from docx import Document
    from src.baa_engine.word_generator import build_correction_notice_docx

    items = [{"function_id": "DIM-001", "entity_id": "e1", "note": "净高不足"}]
    meta = {"project_name": "测试项目", "drawing_name": "test.dxf", "reviewer": "张三"}
    out = build_correction_notice_docx(items, meta, str(tmp_path / "notice.docx"))
    doc = Document(out)
    assert len(doc.paragraphs) > 0
    assert len(doc.tables) > 0
    # 检查标题
    assert doc.paragraphs[0].text == "整改通知单"


def test_unknown_template_fallback(tmp_path):
    from src.baa_engine.word_generator import build_correction_notice_docx

    items = [{"function_id": "DIM-001", "entity_id": "e1", "note": "x"}]
    meta = {"project_name": "P", "drawing_name": "d.dxf", "reviewer": "R"}
    out = build_correction_notice_docx(items, meta, str(tmp_path / "n.docx"), template_code="nonexistent")
    assert Path(out).is_file()


def test_sdk_correction_methods_exist():
    from src.sdk import BAAClient

    c = BAAClient(api_key="test")
    assert hasattr(c, "generate_correction_notice")
    assert hasattr(c, "generate_recheck_opinion")
